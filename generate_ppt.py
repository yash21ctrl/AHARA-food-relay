import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Brand Colors
    C_BG_DARK = RGBColor(11, 28, 24)       # Deep Forest Night #0B1C18
    C_EMERALD = RGBColor(5, 150, 105)      # Vibrant Emerald #059669
    C_TEAL = RGBColor(13, 148, 136)        # Teal #0D9488
    C_AMBER = RGBColor(217, 119, 6)        # Warm Amber #D97706
    C_WHITE = RGBColor(255, 255, 255)
    C_TEXT_LIGHT = RGBColor(226, 232, 240) # Slate 200
    C_MUTED = RGBColor(148, 163, 184)      # Slate 400
    C_CARD_BG = RGBColor(18, 38, 33)       # Dark Card BG
    C_CARD_BORDER = RGBColor(30, 60, 52)   # Border

    def add_bg(slide, dark=True):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG_DARK if dark else RGBColor(248, 250, 252)
        bg.line.fill.background()
        return bg

    def add_header(slide, tag, title, subtitle):
        # Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.5), Inches(3.2), Inches(0.38))
        badge.fill.solid()
        badge.fill.fore_color.rgb = C_CARD_BG
        badge.line.color.rgb = C_EMERALD
        tf_b = badge.text_frame
        tf_b.word_wrap = True
        p_b = tf_b.paragraphs[0]
        p_b.text = f"● {tag.upper()}"
        p_b.font.size = Pt(11)
        p_b.font.bold = True
        p_b.font.color.rgb = C_EMERALD
        p_b.alignment = PP_ALIGN.CENTER

        # Title & Subtitle in single text box to prevent overlap
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.95), Inches(11.7), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(26)
        p1.font.bold = True
        p1.font.color.rgb = C_WHITE

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(13)
        p2.font.color.rgb = C_MUTED
        p2.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, dark=True)

    # Accent decorative box
    glow = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.0), Inches(11.733), Inches(5.5))
    glow.fill.solid()
    glow.fill.fore_color.rgb = C_CARD_BG
    glow.line.color.rgb = C_CARD_BORDER

    # Main Title Box
    t_box = s1.shapes.add_textbox(Inches(1.3), Inches(1.4), Inches(10.7), Inches(3.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "AHARA (FoodRelay)"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p_sub = tf1.add_paragraph()
    p_sub.text = "Hyperlocal Surplus Food Redistribution & Zero-Waste Ecosystem"
    p_sub.font.size = Pt(20)
    p_sub.font.bold = True
    p_sub.font.color.rgb = C_EMERALD
    p_sub.space_before = Pt(8)

    p_desc = tf1.add_paragraph()
    p_desc.text = "Empowering commercial food donors (weddings, caterers, hotels, cafeterias) with real-time FSSAI safe-limit tracking, dynamic fair delivery quoting, and an emergency zero-waste fallback to animal shelters and Gaushalas."
    p_desc.font.size = Pt(14)
    p_desc.font.color.rgb = C_TEXT_LIGHT
    p_desc.space_before = Pt(14)

    # Team Members Box
    team_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(4.8), Inches(10.7), Inches(1.3))
    team_box.fill.solid()
    team_box.fill.fore_color.rgb = RGBColor(12, 45, 38)
    team_box.line.color.rgb = C_EMERALD

    tf_team = team_box.text_frame
    tf_team.word_wrap = True
    p_tlabel = tf_team.paragraphs[0]
    p_tlabel.text = "PROJECT BUILDERS & PARTICIPANTS"
    p_tlabel.font.size = Pt(11)
    p_tlabel.font.bold = True
    p_tlabel.font.color.rgb = C_AMBER

    p_tmembers = tf_team.add_paragraph()
    p_tmembers.text = "Sohan D   •   Yashwanth M   •   Harshith Y A"
    p_tmembers.font.size = Pt(18)
    p_tmembers.font.bold = True
    p_tmembers.font.color.rgb = C_WHITE
    p_tmembers.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 2: THE CRITICAL PROBLEM
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "The Real-World Challenge", "The Paradox of Massive Surplus Food Wastage", "Over 68 million tons of food is wasted annually in India, while millions remain food-insecure.")

    cards2 = [
        ("1. Hyperlocal Information Gap", "Donors (weddings, hotels, caterers) have high-quality freshly cooked food, but lack immediate live visibility into which verified shelter or orphanage within a 5-10 km radius can receive it tonight.", C_EMERALD),
        ("2. FSSAI Food Safety & Spoilage Fear", "Cooked food degrades rapidly. Without clear, immutable preparation-to-consumption safe window limits, donors fear liability, causing good surplus food to be discarded preemptively.", C_AMBER),
        ("3. The Delivery Logistics Deadlock", "Who delivers the food? Volunteers lack vehicles; NGOs operate on shoestring budgets; donors don't know the fare. The lack of fair split-cost delivery mechanisms causes rescue attempts to fail.", C_TEAL),
        ("4. Zero Fallback for Late-Hour Food", "When food nears the end of its consumption window with no human takers, 100% of it is thrown into municipal dumpsters, rotting in landfills and releasing potent methane gas.", RGBColor(225, 29, 72))
    ]

    for i, (title, desc, col) in enumerate(cards2):
        x = Inches(0.8 + (i % 2) * 5.95)
        y = Inches(2.2 + (i // 2) * 2.4)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.15))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = col

        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_LIGHT
        p2.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 3: THE AHARA SOLUTION
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "Comprehensive Innovation", "AHARA: The 360° Food Rescue Ecosystem", "A purpose-built web platform engineered for extreme speed, strict food safety, and zero food wastage.")

    pillars = [
        ("⚡ 60-Second Posting Flow", "Donors publish surplus meal listings in under 1 minute with event type, batch quantity, preparation hour, and storage conditions.", C_EMERALD),
        ("🛡️ Immutable Safe-Until Limits", "FSSAI safety standards embedded into continuous countdown timers that strictly block unsafe expired food distribution.", C_TEAL),
        ("🤝 Dynamic Fare Quoting", "Real-time delivery cost estimation with custom Porter/Dunzo/Uber quoting and 5 flexible cost-sharing negotiation modes.", C_AMBER),
        ("🐾 Zero-Waste Animal Fallback", "Automated 30-minute pre-expiry diversion routing remaining surplus to local cattle shelters (Gaushalas) and animal rescue hubs.", C_EMERALD)
    ]

    for i, (title, desc, col) in enumerate(pillars):
        x = Inches(0.8 + i * 2.95)
        y = Inches(2.2)
        c = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.85), Inches(4.5))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = col

        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(17)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_LIGHT
        p2.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 4: DUAL-PORTAL ARCHITECTURE
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "System Design & Portals", "Multi-Portal Architecture with Multi-Tab Session Isolation", "Tailored user journeys for donors and recipients, supporting simultaneous isolated operations.")

    # Portal 1: Donor
    c_don = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(5.75), Inches(4.5))
    c_don.fill.solid()
    c_don.fill.fore_color.rgb = C_CARD_BG
    c_don.line.color.rgb = C_EMERALD
    tf = c_don.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🍲 DONOR PORTAL (/#/donor)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    items_don = [
        "• Rapid Listing Creation: Specify surplus plates, event context, and storage (Ambient / Chilled / Heated).",
        "• Live Safe-Life Monitoring: Color-coded remaining safe duration with automated alert triggers.",
        "• Delivery Proposal Review: Accept or counter delivery method and custom delivery fare quotes.",
        "• Vehicle & Driver Dispatch: Track driver name, vehicle number plate, phone, and live ETA.",
        "• Emergency Animal Fallback: Divert last-half-hour food to Gaushalas with 1 click."
    ]
    for it in items_don:
        pi = tf.add_paragraph()
        pi.text = it
        pi.font.size = Pt(12)
        pi.font.color.rgb = C_TEXT_LIGHT
        pi.space_before = Pt(6)

    # Portal 2: Receiver
    c_rec = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.75), Inches(2.2), Inches(5.75), Inches(4.5))
    c_rec.fill.solid()
    c_rec.fill.fore_color.rgb = C_CARD_BG
    c_rec.line.color.rgb = C_TEAL
    tf = c_rec.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🏢 RECIPIENT PORTAL (/#/receiver)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    items_rec = [
        "• Localized Live Feed: Explore verified surplus listings across Bangalore urban micro-clusters.",
        "• Flexible Batch Acceptance: Claim partial quantities (e.g. 40 of 100 plates) matching shelter capacity.",
        "• Delivery Negotiation Engine: Propose Self Pickup, Donor Drop, or Shared Commercial Courier.",
        "• Real-Time Status Lifecycle: ACCEPTED → CONFIRMED → READY → OUT_FOR_DELIVERY → DELIVERED.",
        "• Delivery Receipt Verification: Confirm receipt and record community nutritional impact."
    ]
    for it in items_rec:
        pi = tf.add_paragraph()
        pi.text = it
        pi.font.size = Pt(12)
        pi.font.color.rgb = C_TEXT_LIGHT
        pi.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 5: FOOD SAFETY & FSSAI COMPLIANCE
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Safety & Compliance", "FSSAI Food Safety Standards & Immutable Safe Timers", "Public health is paramount: AHARA embeds food safety science directly into code.")

    f_cards = [
        ("🔒 Immutable 'safeUntil' Lock", "Calculated strictly at listing creation as createdAt + safeHours * 3,600,000 ms. Once set, safeUntil can NEVER be extended or tampered with by any party.", C_EMERALD),
        ("⏱️ Live Continuous Countdown", "Active countdown timers on every card and detail screen display exact hours, minutes, and seconds remaining with visual warning states (Green → Amber → Red).", C_TEAL),
        ("🚫 Automated Expiry Lockout", "As soon as safeUntil is reached, human claim actions are disabled instantly. The post transitions to EXPIRED, preventing any distribution of spoiled food.", C_AMBER),
        ("🔔 30-Minute Advance Alert", "At 30 minutes remaining, automated alerts warn both donor and active recipients, allowing orderly handover or diversion before the window closes.", RGBColor(56, 189, 248))
    ]

    for i, (title, desc, col) in enumerate(f_cards):
        x = Inches(0.8 + (i % 2) * 5.95)
        y = Inches(2.2 + (i // 2) * 2.4)
        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.15))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = col

        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_LIGHT
        p2.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 6: DYNAMIC DELIVERY FARE NEGOTIATION
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Logistics Innovation", "Dynamic Fare Quoting & 5-Way Cost Sharing", "Eliminating delivery friction through transparent negotiation and live price quoting.")

    # Left: The 5 Methods
    c_m = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(6.0), Inches(4.6))
    c_m.fill.solid()
    c_m.fill.fore_color.rgb = C_CARD_BG
    c_m.line.color.rgb = C_AMBER
    tf = c_m.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🤝 5 FLEXIBLE DELIVERY METHODS"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_AMBER

    methods = [
        ("• SELF_PICKUP (₹0)", "Recipient shelter arranges their own volunteer or vehicle."),
        ("• DONOR_DROP (₹0)", "Donor drops food off at the shelter on their way."),
        ("• SPLIT (50 / 50)", "Delivery cost is shared equally between donor & recipient shelter."),
        ("• DONOR_PAYS (100%)", "Donor sponsors full delivery courier as an act of charity."),
        ("• RECEIVER_PAYS (100%)", "Recipient NGO covers the delivery fare via institutional funds.")
    ]
    for title, desc in methods:
        pi = tf.add_paragraph()
        pi.text = f"{title}: {desc}"
        pi.font.size = Pt(11)
        pi.font.color.rgb = C_TEXT_LIGHT
        pi.space_before = Pt(5)

    # Right: Live Fare Quoting
    c_q = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(2.2), Inches(5.5), Inches(4.6))
    c_q.fill.solid()
    c_q.fill.fore_color.rgb = C_CARD_BG
    c_q.line.color.rgb = C_EMERALD
    tf = c_q.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🛵 LIVE APP FARE QUOTING (NEW)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    q_points = [
        "1. Haversine Baseline Distance: System computes exact km between Bangalore zones (Indiranagar, Koramangala, HSR, Whitefield, etc.).",
        "2. Live Courier Price Entry: Proposer can check Porter, Rapido, Dunzo, or Uber Auto and enter the exact live quote.",
        "3. Dynamic Share Splitting: Shares update automatically in real time (e.g. ₹160 fare → ₹80 Donor / ₹80 Receiver).",
        "4. Mid-Delivery Adjustments: If driver encounters highway toll or weather surge, costs can be updated during OUT_FOR_DELIVERY with full audit trails.",
        "5. Two-Way Agreement: Both parties must explicitly accept before transport details are confirmed."
    ]
    for q in q_points:
        pi = tf.add_paragraph()
        pi.text = q
        pi.font.size = Pt(11)
        pi.font.color.rgb = C_TEXT_LIGHT
        pi.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 7: ZERO-WASTE ANIMAL SHELTER & GAUSHALA FALLBACK
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Game-Changing Zero Waste", "Zero-Waste Fallback: Gaushala & Animal Shelter Diversion", "When human food windows close, animal feeding ensures 100% zero food goes to waste.")

    z_cards = [
        ("🚨 30-Min Emergency Alert", "When a listing enters the final 30 minutes with remaining plates, donor gets an alert to divert unclaimed food to nearby Gaushalas or animal rescue centers.", C_AMBER),
        ("📍 Nearest Gaushala Discovery", "Curated registry of verified Bangalore animal shelters (Sri Krishna Gaushala, CARE, Akhila Karnataka Gaushala) automatically sorted by distance with phone contacts.", C_EMERALD),
        ("❤️ Donor Goodwill Courier", "Donor can deliver directly (Self-Drop ₹0) or book a Goodwill Courier where the donor covers the delivery fare (~₹80) as an act of compassion & Gau Seva.", C_TEAL),
        ("🎉 Celebration Pop-Up & Quotes", "Successful dispatch triggers an uplifting celebration modal with inspiring quotes (Gandhi, Buddha, Jane Goodall), marking the rescue 100% zero-waste complete.", RGBColor(244, 114, 182))
    ]

    for i, (title, desc, col) in enumerate(z_cards):
        x = Inches(0.8 + (i % 2) * 5.95)
        y = Inches(2.2 + (i // 2) * 2.4)
        c = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.15))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = col

        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_LIGHT
        p2.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 8: IMPACT DASHBOARD & METRICS
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "Metrics & Accountability", "Real-Time Environmental & Community Impact Dashboard", "Quantifying meals delivered, hunger relieved, and carbon emissions averted in real time.")

    stats = [
        ("Plates Saved", "Human + Animal Rescues Combined", C_EMERALD),
        ("Verified Orders", "100% Receipt Confirmed by Shelters", C_TEAL),
        ("Animals Fed", "Dispatched to Gaushalas & Rescue Hubs", C_AMBER),
        ("CO₂e Prevented", "Organic Rot Landfill Mitigation", RGBColor(56, 189, 248))
    ]

    for i, (title, desc, col) in enumerate(stats):
        x = Inches(0.8 + i * 2.95)
        y = Inches(2.2)
        c = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(2.85), Inches(2.0))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = col

        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = C_TEXT_LIGHT
        p2.space_before = Pt(6)

    # Environmental Equivalents Banner Box
    banner = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.5), Inches(11.733), Inches(2.3))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(12, 45, 38)
    banner.line.color.rgb = C_EMERALD
    tf_b = banner.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "🌿 SCIENTIFIC IMPACT MULTIPLIERS EMBEDDED"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    items_b = [
        "• Food Mass Equivalent: ~400 grams per meal plate rescued from going into municipal landfills.",
        "• Greenhouse Gas Mitigation: ~1.5 kg CO₂e greenhouse gas averted per kg of surplus food rescued.",
        "• Economic Value Generated: ₹75 per human meal plate & ₹40 per cattle feed plate delivered directly to grassroot charities.",
        "• Dual Audit Tables: Real-time logs for both verified shelter deliveries and Gaushala zero-waste dispatches."
    ]
    for b in items_b:
        pi = tf_b.add_paragraph()
        pi.text = b
        pi.font.size = Pt(12)
        pi.font.color.rgb = C_WHITE
        pi.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 9: SYSTEM ROBUSTNESS & CONCURRENCY
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9)
    add_header(s9, "Engineering Excellence", "Concurrency, Edge-Case Handling & Automated Testing", "Built for real-world stress conditions with defensive programming and zero silent errors.")

    r_cards = [
        ("⚡ Simultaneous Claim Protection", "Atomic inventory checks prevent race conditions when multiple recipients claim simultaneously. If remainingQty is depleted, the second recipient receives an instant friendly block with no data corruption.", C_EMERALD),
        ("🔄 Cancellation Inventory Reversal", "When an order is cancelled, plates are immediately refunded back to the post's active remainingQty and reopened for nearby shelters, broadcasting restock alerts across Bangalore.", C_AMBER),
        ("📝 Mandatory Reason & Audit Trail", "Cancellations strictly require a verified reason (min 5 chars), recording who cancelled and why, displayed transparently in order histories to deter frivolous cancellations.", C_TEAL),
        ("🧪 8-Test Automated Test Suite", "Full test-suite.js verifies Haversine distance, custom fare splits, zero-fare promos, mid-delivery adjustments, time offset simulation, and Gaushala 30-min dispatches.", RGBColor(168, 85, 247))
    ]

    for i, (title, desc, col) in enumerate(r_cards):
        x = Inches(0.8 + (i % 2) * 5.95)
        y = Inches(2.2 + (i // 2) * 2.4)
        c = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.15))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = col

        tf = c.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_LIGHT
        p2.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 10: FUTURE ROADMAP & TEAM SUMMARY
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10)
    add_header(s10, "Looking Ahead", "Scalability Roadmap & Team Credits", "Scaling AHARA to eliminate hunger and food wastage across Indian urban hubs.")

    # Roadmap
    c_road = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(6.0), Inches(4.6))
    c_road.fill.solid()
    c_road.fill.fore_color.rgb = C_CARD_BG
    c_road.line.color.rgb = C_EMERALD
    tf = c_road.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🚀 SCALABILITY ROADMAP"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD

    r_items = [
        "1. IoT Cold-Chain Temperature Sensors: Bluetooth/NFC sensors logging container temperatures for milk/paneer dishes.",
        "2. WhatsApp & IVR Voice Bot: Enabling grassroots shelter caretakers to accept food without requiring smartphones.",
        "3. Automated Tax Receipts (Section 80G): Instant digital donation deduction certificates for commercial food caterers.",
        "4. Electric Vehicle Delivery Fleet Partnership: Partnering with green EV logistics fleets for zero-emission food transport."
    ]
    for it in r_items:
        pi = tf.add_paragraph()
        pi.text = it
        pi.font.size = Pt(11)
        pi.font.color.rgb = C_TEXT_LIGHT
        pi.space_before = Pt(7)

    # Team Credits
    c_team = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(2.2), Inches(5.5), Inches(4.6))
    c_team.fill.solid()
    c_team.fill.fore_color.rgb = RGBColor(12, 45, 38)
    c_team.line.color.rgb = C_AMBER
    tf = c_team.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👥 PROJECT CREATORS"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_AMBER

    members = [
        ("Sohan D", "Fullstack Architecture, State Management & System Design"),
        ("Yashwanth M", "Frontend Engineering, Dynamic Delivery Quoting & Portals"),
        ("Harshith Y A", "Zero-Waste Logistics, Impact Metrics & Quality Engineering")
    ]
    for name, role in members:
        pi = tf.add_paragraph()
        pi.text = f"⭐ {name}"
        pi.font.size = Pt(15)
        pi.font.bold = True
        pi.font.color.rgb = C_WHITE
        pi.space_before = Pt(8)

        pr = tf.add_paragraph()
        pr.text = role
        pr.font.size = Pt(11)
        pr.font.color.rgb = C_TEXT_LIGHT

    p_end = tf.add_paragraph()
    p_end.text = "\nThank you! Questions & Live Demo Welcome."
    p_end.font.size = Pt(13)
    p_end.font.bold = True
    p_end.font.color.rgb = C_EMERALD

    output_path = os.path.join(os.getcwd(), "AHARA_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_deck()
